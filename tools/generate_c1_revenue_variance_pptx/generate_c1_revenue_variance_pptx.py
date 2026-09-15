"""
generate_c1_revenue_variance_pptx.py
=====================================
watsonx Orchestrate tool — converts a pre-assembled slide data JSON payload
into a filled PPTX variance-analysis deck using python-pptx.

All IBM Planning Analytics data fetching is done OUTSIDE this tool by the
calling agent (pa_agent) using the ibm-pa-tools MCP toolkit.
This tool is purely a "data in → .pptx bytes out" renderer.

Import command
--------------
    orchestrate tools import \\
        -k python \\
        -f tools/generate_c1_revenue_variance_pptx/generate_c1_revenue_variance_pptx.py \\
        -r tools/generate_c1_revenue_variance_pptx/requirements.txt \\
        -p tools/generate_c1_revenue_variance_pptx
"""

import io
import json
import os
from datetime import date
from lxml import etree

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Pt

from ibm_watsonx_orchestrate.agent_builder.tools import ToolPermission, tool

# Absolute path to the deployed package directory (variance_template.pptx lives here)
_TOOL_DIR = os.path.dirname(os.path.abspath(__file__))
_TEMPLATE  = os.path.join(_TOOL_DIR, "variance_template.pptx")

# ── Colour constants ──────────────────────────────────────────────────────────
_GREEN   = "28A745"
_RED     = "DC3545"
_NEUTRAL = "0F2040"
_BLUE    = "1E90FF"
_MUTED   = "A0B8D8"


# ── Text helpers ──────────────────────────────────────────────────────────────

def _copy_rpr(src_run, dst_run):
    """Copy font properties (size, name, bold, color) from src run to dst run."""
    sf = src_run.font
    df = dst_run.font
    if sf.size:
        df.size = sf.size
    if sf.name:
        df.name = sf.name
    if sf.bold is not None:
        df.bold = sf.bold
    # Copy color if it's an explicit RGB
    try:
        rgb = sf.color.rgb
        df.color.rgb = rgb
    except Exception:
        pass


def _set_text(shape, text: str, color_hex: str | None = None,
              bold: bool | None = None, font_size: float | None = None):
    """Replace all text in a shape with *text*, preserving the template's font
    properties (size, name) and paragraph spacing/alignment.

    Only override color/bold/font_size when explicitly passed.
    """
    tf = shape.text_frame
    # Snapshot what the first paragraph/run looks like before we clear
    first_para = tf.paragraphs[0] if tf.paragraphs else None
    orig_align = first_para.alignment if first_para else None
    # Snapshot spacing from the paragraph XML so we can restore it
    orig_para_xml = etree.tostring(first_para._p) if first_para is not None else None
    # Snapshot the first run's font properties
    first_run = first_para.runs[0] if (first_para and first_para.runs) else None

    tf.clear()
    p = tf.paragraphs[0]

    # Restore paragraph-level spacing/alignment from original XML
    if orig_para_xml is not None:
        orig_p = etree.fromstring(orig_para_xml)
        # Copy <a:pPr> from original paragraph if it exists
        orig_pPr = orig_p.find(qn("a:pPr"))
        if orig_pPr is not None:
            existing_pPr = p._p.find(qn("a:pPr"))
            if existing_pPr is not None:
                p._p.remove(existing_pPr)
            p._p.insert(0, etree.fromstring(etree.tostring(orig_pPr)))

    if orig_align and p.alignment is None:
        p.alignment = orig_align

    run = p.add_run()
    run.text = text

    # Apply template font properties first (size, name from original run)
    if first_run is not None:
        _copy_rpr(first_run, run)

    # Then apply explicit overrides
    if color_hex:
        run.font.color.rgb = RGBColor.from_string(color_hex)
    if bold is not None:
        run.font.bold = bold
    if font_size:
        run.font.size = Pt(font_size)


def _set_text_multi(shape, paragraphs: list[dict]):
    """Replace shape text with multiple paragraphs.

    Each dict: {"text": str, "color_hex": str|None, "bold": bool|None}
    Font size / name / spacing is copied from the corresponding original paragraph.
    """
    tf = shape.text_frame
    orig_paras = list(tf.paragraphs)

    tf.clear()
    for i, spec in enumerate(paragraphs):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        # Restore paragraph spacing from original para at same index (or last if shorter)
        orig_idx = min(i, len(orig_paras) - 1)
        if orig_paras:
            orig_p_xml = etree.tostring(orig_paras[orig_idx]._p)
            orig_p_elem = etree.fromstring(orig_p_xml)
            orig_pPr = orig_p_elem.find(qn("a:pPr"))
            if orig_pPr is not None:
                existing_pPr = p._p.find(qn("a:pPr"))
                if existing_pPr is not None:
                    p._p.remove(existing_pPr)
                p._p.insert(0, etree.fromstring(etree.tostring(orig_pPr)))

        run = p.add_run()
        run.text = spec["text"]

        # Copy font from original run at same para index
        if orig_paras and orig_idx < len(orig_paras):
            orig_runs = orig_paras[orig_idx].runs
            if orig_runs:
                _copy_rpr(orig_runs[0], run)

        if spec.get("color_hex"):
            run.font.color.rgb = RGBColor.from_string(spec["color_hex"])
        if spec.get("bold") is not None:
            run.font.bold = spec["bold"]


# ── Numeric helpers ───────────────────────────────────────────────────────────

def _to_float(s) -> float:
    """Parse a value like '+78.9%' / '-0.9%' / '1234567' to float."""
    s = str(s).replace(",", "").replace("(", "-").replace(")", "").strip().rstrip("%").replace("+", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _fmt_pct(n: float) -> str:
    sign = "+" if n > 0 else ""
    return f"{sign}{n:.1f}%"


def _var_color(positive: bool | None = None, var_pct: float | None = None) -> str:
    if positive is not None:
        return _GREEN if positive else _RED
    if var_pct is not None:
        return _GREEN if var_pct >= 0 else _RED
    return _NEUTRAL


# ── Auto-derive findings ──────────────────────────────────────────────────────

def _build_findings(kpi: dict, breakdowns: list, point_a: str, point_b: str):
    total_a  = kpi["total_a"] / 1e6
    total_b  = kpi["total_b"] / 1e6
    var_pct  = kpi["var_pct"]
    overall_color = _GREEN if var_pct >= 0 else _RED
    overall_tag   = "STRONG GROWTH" if var_pct >= 10 else ("DECLINE" if var_pct < 0 else "STABLE")
    findings = [{
        "tag":   overall_tag,
        "title": f"Overall Revenue {_fmt_pct(var_pct)} YoY",
        "body":  f"Total grew from ${total_a:.1f}M ({point_a}) to ${total_b:.1f}M ({point_b}).",
        "color": overall_color,
    }]

    all_stats = [(s, bd["dim_name"]) for bd in breakdowns for s in bd.get("stats", [])]
    if all_stats:
        best  = max(all_stats, key=lambda x: _to_float(x[0]["var"]))
        worst = min(all_stats, key=lambda x: _to_float(x[0]["var"]))
        findings.append({
            "tag":   "TOP MOVER",
            "title": f"{best[0]['name']}: {best[0]['var']}",
            "body":  f"Strongest performer across all {best[1]}. {best[0]['val_a']} → {best[0]['val_b']}.",
            "color": "FF8C00",
        })
        if worst[0] is not best[0]:
            findings.append({
                "tag":   "LAGGARD",
                "title": f"{worst[0]['name']}: {worst[0]['var']}",
                "body":  f"Weakest performer. {worst[0]['val_a']} → {worst[0]['val_b']}.",
                "color": _RED,
            })

    for bd in breakdowns[:2]:
        if len(findings) >= 5:
            break
        if bd.get("stats"):
            top = max(bd["stats"], key=lambda s: _to_float(s["var"]))
            findings.append({
                "tag":   bd["dim_name"].upper()[:12],
                "title": f"{top['name']}: {top['var']}",
                "body":  f"Top {top['name']} in {bd['dim_name']}: {top['val_a']} → {top['val_b']}.",
                "color": _BLUE,
            })

    while len(findings) < 5:
        findings.append({"tag": "", "title": "", "body": "", "color": "888888"})

    recs = [
        f"Investigate drivers behind {findings[1]['title'].split(':')[0].strip()}" if len(findings) > 1 else "Analyze top movers",
        f"Address underperformance in {findings[2]['title'].split(':')[0].strip()}" if len(findings) > 2 else "Monitor laggards",
        "Review seasonal patterns and plan for low-revenue periods",
        "Align targets for next period based on observed variance",
    ]
    return findings[:5], recs


# ── Slide fillers ─────────────────────────────────────────────────────────────

def _fill_slide0(slide, meta):
    """Title slide.

    Document-order shape indices (from slide.shapes[i]):
      s[0]: background rect (no text frame, skip)
      s[1]: eyebrow         "IBM PLANNING ANALYTICS · …"  ← leave static
      s[2]: empty text box  (leave)
      s[3]: main title      "C1_Revenue\x0bVariance Analysis"
      s[4]: subtitle        "2024 vs. 2025 Actual Gross Revenue"
      s[5]: context line    "Total Company · All Channels…"
    """
    s = slide.shapes

    # s[3]: title — uses a single paragraph with <a:br/> (soft return) in template.
    # We edit in-place: update the first run's text to the cube name,
    # then update (or add) the run after the <a:br/> to "Variance Analysis".
    cube_name = meta.get("cube", "C1_Revenue")
    tf_title = s[3].text_frame
    para = tf_title.paragraphs[0]

    # Collect all child elements of <a:p>
    children = list(para._p)
    runs_in_para = [c for c in children if c.tag == qn("a:r")]
    brs_in_para  = [c for c in children if c.tag == qn("a:br")]

    if runs_in_para:
        # Update first run text to cube_name
        runs_in_para[0].find(qn("a:t")).text = cube_name
        # Ensure white bold color on first run
        _apply_rpr_overrides(runs_in_para[0], color_hex="FFFFFF", bold=True)

    if len(runs_in_para) >= 2:
        # Update second run (after <a:br/>) to "Variance Analysis"
        runs_in_para[1].find(qn("a:t")).text = "Variance Analysis"
        _apply_rpr_overrides(runs_in_para[1], color_hex="FFFFFF", bold=True)
    elif brs_in_para:
        # <a:br/> exists but no second run — add one
        r2_elem = etree.SubElement(para._p, qn("a:r"))
        t2 = etree.SubElement(r2_elem, qn("a:t"))
        t2.text = "Variance Analysis"
        if runs_in_para:
            rPr_src = runs_in_para[0].find(qn("a:rPr"))
            if rPr_src is not None:
                r2_elem.insert(0, etree.fromstring(etree.tostring(rPr_src)))
        _apply_rpr_overrides(r2_elem, color_hex="FFFFFF", bold=True)
    else:
        # No <a:br/> at all — just set the single run
        if runs_in_para:
            runs_in_para[0].find(qn("a:t")).text = f"{cube_name}\x0bVariance Analysis"

    # s[4]: subtitle
    _set_text(s[4],
              f"{meta.get('point_a','A')} vs. {meta.get('point_b','B')} "
              f"Actual {meta.get('measure','Gross Revenue')}",
              color_hex=_MUTED)

    # s[5]: context line
    _set_text(s[5], meta.get("subtitle", ""), color_hex="6B8CAE")


def _apply_rpr_overrides(r_elem, color_hex: str | None = None, bold: bool | None = None):
    """Apply color and bold overrides to an <a:r> element's <a:rPr>.

    Inserts <a:solidFill> before <a:latin> so the OOXML child order is valid —
    PowerPoint ignores fill elements that appear after typeface declarations.
    """
    rPr = r_elem.find(qn("a:rPr"))
    if rPr is None:
        rPr = etree.Element(qn("a:rPr"))
        r_elem.insert(0, rPr)
    if bold is not None:
        rPr.set("b", "1" if bold else "0")
    if color_hex:
        # Remove any existing fill elements
        for fill in rPr.findall(qn("a:solidFill")):
            rPr.remove(fill)
        fill_elem = etree.Element(qn("a:solidFill"))
        clr_elem  = etree.SubElement(fill_elem, qn("a:srgbClr"))
        clr_elem.set("val", color_hex.upper())
        # Insert before <a:latin> so element order matches OOXML spec
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addprevious(fill_elem)
        else:
            rPr.insert(0, fill_elem)


def _set_xfrm(shape, top_emu: int | None = None, left_emu: int | None = None,
              cx_emu: int | None = None, cy_emu: int | None = None):
    """Move / resize a shape by updating its <a:xfrm> <a:off> and <a:ext>."""
    xfrm = shape._element.find(".//" + qn("a:xfrm"))
    if xfrm is None:
        return
    off = xfrm.find(qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    if off is not None:
        if left_emu is not None:
            off.set("x", str(left_emu))
        if top_emu is not None:
            off.set("y", str(top_emu))
    if ext is not None:
        if cx_emu is not None:
            ext.set("cx", str(cx_emu))
        if cy_emu is not None:
            ext.set("cy", str(cy_emu))


def _fill_slide1(slide, meta, kpi):
    """Executive KPI summary.

    Document-order shape indices:
      s[0]:  background rect (no text)
      s[1]:  decorative shape (no text)
      s[2]:  slide title       "Executive Summary — Revenue Overview"
      s[3]:  subtitle          "C1_Revenue Cube · …"
      s[4]:  KPI box A bg (no text)
      s[5]:  KPI box A border (no text)
      s[6]:  KPI label A       "2024 Full Year / Gross Revenue"   ← update
      s[7]:  KPI val A         "$211.9M"
      s[8]:  KPI badge A       "Baseline Year"                    ← static
      s[9]:  KPI box B bg (no text)
      s[10]: KPI box B border (no text)
      s[11]: KPI label B       "2025 Full Year / Gross Revenue"   ← update
      s[12]: KPI val B         "$284.8M"
      s[13]: % change badge    "▲ +34.4%"
      s[14]: KPI box Δ bg (no text)
      s[15]: KPI box Δ border (no text)
      s[16]: KPI label Δ       "Year-over-Year / Variance ($)"    ← static
      s[17]: KPI val Δ         "+$72.9M"
      s[18]: KPI badge Δ       "Absolute Growth"                  ← static
      s[19]: KPI box Q bg (no text)
      s[20]: KPI box Q border (no text)
      s[21]: KPI label Q       "Best Performing / Quarter"        ← static
      s[22]: KPI val Q         "Q4 2025"
      s[23]: KPI detail Q      "$83.1M (+30.6%)"
    """
    s = slide.shapes
    pa      = meta.get("point_a", "A")
    pb      = meta.get("point_b", "B")
    measure = meta.get("measure", "Gross Revenue")
    var_col = _GREEN if kpi["var_pct"] >= 0 else _RED

    # s[2]: slide title
    _set_text(s[2], "Executive Summary — Revenue Overview", color_hex="FFFFFF", bold=True)

    # s[3]: subtitle
    _set_text(s[3],
              f"{meta.get('cube','C1_Revenue')} Cube · {meta.get('subtitle','Total Company · Actual')}",
              color_hex=_MUTED)

    # s[6], s[11], s[16], s[21]: KPI column label boxes — move up into the top
    # portion of each card so they don't overlap the large value text below.
    # Original top=2516088 EMU (2.75"), new top=1800072 EMU (1.97") with taller box.
    _LABEL_TOP = 1800072   # ~1.97" — sits just below the card header line
    _LABEL_H   = 457200    # ~0.50" — enough for two 10pt lines with spacing
    for label_shape_idx in (6, 11, 16, 21):
        _set_xfrm(s[label_shape_idx], top_emu=_LABEL_TOP, cy_emu=_LABEL_H)

    # s[6]: label for KPI card A (2 paragraphs)
    _set_text_multi(s[6], [
        {"text": f"{pa} Full Year", "color_hex": "6B8CAE"},
        {"text": measure,           "color_hex": "6B8CAE"},
    ])

    # s[7]: total_a value
    _set_text(s[7], f"${kpi['total_a']/1e6:.1f}M", color_hex=_NEUTRAL, bold=True)

    # s[8]: "Baseline Year" ← static, leave

    # s[11]: label for KPI card B (2 paragraphs)
    _set_text_multi(s[11], [
        {"text": f"{pb} Full Year", "color_hex": "6B8CAE"},
        {"text": measure,           "color_hex": "6B8CAE"},
    ])

    # s[12]: total_b value
    _set_text(s[12], f"${kpi['total_b']/1e6:.1f}M", color_hex=var_col, bold=True)

    # s[13]: % change badge
    arrow = "▲" if kpi["var_pct"] >= 0 else "▼"
    _set_text(s[13], f"{arrow} {_fmt_pct(kpi['var_pct'])}", color_hex=var_col, bold=True)

    # s[16]: "Year-over-Year / Variance ($)" ← static, leave

    # s[17]: variance absolute
    sign = "+" if kpi["var_abs"] >= 0 else ""
    _set_text(s[17], f"{sign}${abs(kpi['var_abs'])/1e6:.1f}M", color_hex=var_col, bold=True)

    # s[18]: "Absolute Growth" ← static, leave

    # s[21]: "Best Performing / Quarter" ← static, leave

    # s[22]: best period value
    best = kpi.get("best_period", "")
    _set_text(s[22], f"{best} {pb}", color_hex=_NEUTRAL, bold=True)

    # s[23]: best period detail
    bval = kpi.get("best_period_val_b")
    bpct = kpi.get("best_period_pct")
    if bval is not None and bpct is not None:
        _set_text(s[23], f"${bval/1e6:.1f}M ({_fmt_pct(bpct)})", color_hex=_BLUE, bold=True)


def _fill_slide2(slide, trend):
    """Monthly trend: update text cards + line chart.

    Document-order shape indices:
      s[0]:  background rect (no text)
      s[1]:  decorative shape (no text)
      s[2]:  slide title       "Monthly Revenue Trend — …"
      s[3]:  subtitle          "Gross Revenue by Month…"  ← static, leave
      s[4]:  card-0 bg (no text)
      s[5]:  card-0 border (no text)
      s[6]:  card-0 empty header area (no useful text)
      s[7]:  card-0 tag        "GROWTH"                   ← write insight[0].tag
      s[8]:  card-0 title      "Consistent YoY Uplift"    ← write insight[0].title
      s[9]:  card-0 body       "Every month…"             ← write insight[0].body
      s[10]: card-1 bg (no text)
      s[11]: card-1 border (no text)
      s[12]: card-1 empty header area
      s[13]: card-1 tag        "PEAK"
      s[14]: card-1 title      "Dec 2025: $35.9M"
      s[15]: card-1 body       "Strongest single month…"
      s[16]: card-2 bg (no text)
      s[17]: card-2 border (no text)
      s[18]: card-2 empty header area
      s[19]: card-2 tag        "DIP"
      s[20]: card-2 title      "August Seasonality"
      s[21]: card-2 body       "Aug weakest…"
      s[22]: chart (no text frame)
    """
    if trend is None:
        return
    s = slide.shapes

    label_a = trend.get("label_a", "Prior Year")
    label_b = trend.get("label_b", "Current Year")

    # s[2]: slide title
    _set_text(s[2], f"Monthly Revenue Trend — {label_a} vs. {label_b}",
              color_hex="FFFFFF", bold=True)

    # s[3]: static subtitle — leave

    insights = trend.get("insights", [{}, {}, {}])

    # Widen and raise all three tag badge boxes so multi-word tags like
    # "PEAK MONTH" render on one line (7pt bold).  Original widths were sized
    # for single short words ("GROWTH"=0.45", "PEAK"=0.28", "DIP"=0.17").
    # New size: 1.60" wide × 0.18" tall, keeping the same left edge.
    _TAG_W  = 1463040   # 1.60" in EMU
    _TAG_H  =  164592   # 0.18" in EMU — allows one tightly-spaced 7pt line
    for tag_idx in (7, 13, 19):
        _set_xfrm(s[tag_idx], cx_emu=_TAG_W, cy_emu=_TAG_H)

    # Card 0: tag=s[7], title=s[8], body=s[9]
    i0 = insights[0] if len(insights) > 0 else {}
    _set_text(s[7], i0.get("tag",   ""), color_hex=_BLUE,    bold=True)
    _set_text(s[8], i0.get("title", ""), color_hex=_NEUTRAL,  bold=True)
    _set_text(s[9], i0.get("body",  ""), color_hex="4A5E72")

    # Card 1: tag=s[13], title=s[14], body=s[15]
    i1 = insights[1] if len(insights) > 1 else {}
    _set_text(s[13], i1.get("tag",   ""), color_hex=_BLUE,    bold=True)
    _set_text(s[14], i1.get("title", ""), color_hex=_NEUTRAL,  bold=True)
    _set_text(s[15], i1.get("body",  ""), color_hex="4A5E72")

    # Card 2: tag=s[19], title=s[20], body=s[21]
    i2 = insights[2] if len(insights) > 2 else {}
    _set_text(s[19], i2.get("tag",   ""), color_hex=_BLUE,   bold=True)
    _set_text(s[20], i2.get("title", ""), color_hex=_NEUTRAL, bold=True)
    _set_text(s[21], i2.get("body",  ""), color_hex="4A5E72")

    # Update embedded line chart (s[22])
    _update_chart(slide, trend["labels"], label_a, trend["values_a"], label_b, trend["values_b"])


def _fill_breakdown_slide(slide, bd):
    """Fill one breakdown slide (slides 3 & 4) with stats panel + bar chart.

    Document-order shape indices (identical for both breakdown slides):
      s[0]:  background rect (no text)
      s[1]:  decorative shape (no text)
      s[2]:  slide title
      s[3]:  subtitle
      s[4]:  row-0 separator line (no text)
      s[5]:  row-0 name        e.g. "PCs (30000)"
      s[6]:  row-0 col label A "2024"
      s[7]:  row-0 val_a       "$97.4M"
      s[8]:  row-0 col label B "2025"
      s[9]:  row-0 val_b       "$174.2M"
      s[10]: row-0 col "Change"← static
      s[11]: row-0 var%        "+78.9%"
      s[12]: row-0 divider (no text)
      s[13]: row-1 separator (no text)
      s[14]: row-1 name
      s[15]: row-1 col label A
      s[16]: row-1 val_a
      s[17]: row-1 col label B
      s[18]: row-1 val_b
      s[19]: row-1 col "Change"← static
      s[20]: row-1 var%
      s[21]: row-1 divider (no text)
      s[22]: row-2 separator (no text)
      s[23]: row-2 name
      s[24]: row-2 col label A
      s[25]: row-2 val_a
      s[26]: row-2 col label B
      s[27]: row-2 val_b
      s[28]: row-2 col "Change"← static
      s[29]: row-2 var%
      s[30]: row-2 divider (no text)
      s[31]: chart
    """
    s = slide.shapes
    label_a = bd.get("label_a", "A")
    label_b = bd.get("label_b", "B")

    # s[2]: title
    _set_text(s[2], bd.get("title", f"Revenue by {bd.get('dim_name','')}"),
              color_hex="FFFFFF", bold=True)
    # s[3]: subtitle
    _set_text(s[3], bd.get("subtitle", ""), color_hex=_MUTED)

    stats = bd.get("stats", [])
    # Each row: (name_i, lbl_a_i, val_a_i, lbl_b_i, val_b_i, var_i)
    # "Change" label shapes (s[10], s[19], s[28]) are left untouched.
    row_offsets = [
        (5,  6,  7,  8,  9,  11),   # row 0
        (14, 15, 16, 17, 18, 20),   # row 1
        (23, 24, 25, 26, 27, 29),   # row 2
    ]
    for row_idx, (ni, lai, vai, lbi, vbi, vri) in enumerate(row_offsets):
        if row_idx < len(stats):
            st  = stats[row_idx]
            pos = st.get("positive", _to_float(st.get("var", 0)) >= 0)
            vc  = _var_color(positive=pos)
            _set_text(s[ni],  st.get("name", ""),   color_hex=_NEUTRAL, bold=True)
            _set_text(s[lai], label_a,               color_hex="6B8CAE")
            _set_text(s[vai], st.get("val_a", ""),   color_hex=_NEUTRAL, bold=True)
            _set_text(s[lbi], label_b,               color_hex="6B8CAE")
            _set_text(s[vbi], st.get("val_b", ""),   color_hex=vc,       bold=True)
            _set_text(s[vri], str(st.get("var", "")), color_hex=vc,      bold=True)
        else:
            # blank out unused rows
            for idx in row_offsets[row_idx]:
                _set_text(s[idx], "")

    # Update bar chart (s[31])
    _update_chart(slide, bd["labels"], label_a, bd["values_a"], label_b, bd["values_b"])


def _fill_slide_findings(slide, meta, findings, recommendations):
    """Key Findings & Recommendations — 2-column grid + recs block.

    Document-order shape indices:
      s[0]:  background rect (no text)
      s[1]:  slide title       "Key Findings & Recommendations"  ← static, leave
      s[2]:  subtitle          "C1_Revenue Analysis Summary…"
      s[3]:  left-col card-0 bg (no text)
      s[4]:  left-col card-0 border (no text)
      s[5]:  finding-0 tag     "STRONG GROWTH"
      s[6]:  finding-0 title   "Overall Revenue Up +34.4%"
      s[7]:  finding-0 body    "Total grew…"
      s[8]:  left-col card-1 bg (no text)
      s[9]:  left-col card-1 border (no text)
      s[10]: finding-2 tag     "STAR PERFORMER"
      s[11]: finding-2 title   "PCs: +78.9% YoY"
      s[12]: finding-2 body    "PC revenue…"
      s[13]: left-col card-2 bg (no text)
      s[14]: left-col card-2 border (no text)
      s[15]: finding-4 tag     "DECLINING"
      s[16]: finding-4 title   "Phones: Flat (-0.9%)"
      s[17]: finding-4 body    "Marginal decline…"
      s[18]: right-col card-0 bg (no text)
      s[19]: right-col card-0 border (no text)
      s[20]: finding-1 tag     "CHANNEL SHIFT"
      s[21]: finding-1 title   "Indirect Channel +94.8%"
      s[22]: finding-1 body    "Fastest-growing…"
      s[23]: right-col card-1 bg (no text)
      s[24]: right-col card-1 border (no text)
      s[25]: finding-3 tag     "SEASONALITY"
      s[26]: finding-3 title   "Q4 & August Patterns"
      s[27]: finding-3 body    "Q4 strongest…"
      s[28]: recs card bg (no text)
      s[29]: recs header       "Recommended Next Steps"
      s[30]: recs body         "Investigate PC demand · …"
    """
    s = slide.shapes
    pa = meta.get("point_a", "A")
    pb = meta.get("point_b", "B")

    # s[2]: subtitle
    _set_text(s[2],
              f"{meta.get('cube','C1_Revenue')} Analysis Summary · {pa} to {pb}",
              color_hex=_MUTED)

    # Grid: (tag_idx, title_idx, body_idx) — findings in display order
    # Left column: findings[0], [2], [4]
    # Right column: findings[1], [3]
    grid = [
        (5,  6,  7),    # finding 0 (left col, top)
        (20, 21, 22),   # finding 1 (right col, top)
        (10, 11, 12),   # finding 2 (left col, middle)
        (25, 26, 27),   # finding 3 (right col, middle)
        (15, 16, 17),   # finding 4 (left col, bottom)
    ]
    for i, (ti, tli, bi) in enumerate(grid):
        f  = findings[i] if i < len(findings) else {"tag": "", "title": "", "body": "", "color": "888888"}
        fc = f.get("color", "888888")
        _set_text(s[ti],  f.get("tag",   ""), color_hex=fc,       bold=True)
        _set_text(s[tli], f.get("title", ""), color_hex="FFFFFF", bold=True)
        _set_text(s[bi],  f.get("body",  ""), color_hex="8AABBF")

    # Recs block — s[29] header, s[30] body
    _set_text(s[29], "Recommended Next Steps", color_hex="FFFFFF", bold=True)
    rec_text = " · ".join(recommendations) if recommendations else ""
    _set_text(s[30], rec_text, color_hex="D0E8FF")


# ── Chart update helper ───────────────────────────────────────────────────────

def _update_chart(slide, labels, label_a, values_a, label_b, values_b):
    """Find the first chart in the slide and replace its data."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            continue
        if not shape.has_chart:
            continue
        chart = shape.chart
        cd = ChartData()
        cd.categories = labels
        cd.add_series(label_a, [float(v) for v in values_a])
        cd.add_series(label_b, [float(v) for v in values_b])
        chart.replace_data(cd)
        return  # only update the first chart per slide


# ── Main @tool ────────────────────────────────────────────────────────────────

@tool(
    name="generate_c1_revenue_variance_pptx",
    description=(
        "Generate a variance analysis PowerPoint deck (.pptx) for the C1_Revenue cube "
        "from a pre-assembled JSON data payload. "
        "The calling agent must supply all Planning Analytics data (KPIs, trend series, "
        "breakdown tables) as a JSON string in the 'slide_data_json' parameter. "
        "Returns the raw .pptx file bytes."
    ),
    permission=ToolPermission.READ_ONLY,
)
def generate_c1_revenue_variance_pptx(
    slide_data_json: str,
    filename: str = "C1_Revenue_variance.pptx",
) -> bytes:
    """
    Args:
        slide_data_json (str): JSON string with the complete slide data payload.
            Required top-level keys:
              meta        – dict: title, cube, server, variance_dimension,
                                  point_a, point_b, label_a, label_b,
                                  measure, subtitle, generated (ISO date)
              kpi         – dict: total_a (float, raw $), total_b, var_abs, var_pct (%)
                                  best_period (str), best_period_val_b (float), best_period_pct (float)
              trend       – dict or null:
                              dim_name, labels (list[str]), values_a (list[float]),
                              values_b (list[float]), label_a, label_b, y_label,
                              insights (list[{tag, title, body}])
              breakdowns  – list[dict], one entry per breakdown dimension:
                              dim_key, dim_name, title, subtitle,
                              labels (list[str]), values_a, values_b,
                              stats (list[{name, val_a, val_b, var, positive}]),
                              label_a, label_b, y_label, bar_dir ("col"|"bar")
            Optional keys:
              findings        – list[dict] (auto-generated if absent)
              recommendations – list[str] (auto-generated if absent)

        filename (str): Desired output filename (informational only; bytes are returned).

    Returns:
        bytes: Raw .pptx file content.
    """
    # ── Parse incoming payload ────────────────────────────────────────────────
    try:
        data = json.loads(slide_data_json)
    except json.JSONDecodeError as exc:
        raise ValueError(f"slide_data_json is not valid JSON: {exc}") from exc

    for key in ("meta", "kpi", "breakdowns"):
        if key not in data:
            raise ValueError(f"slide_data_json is missing required key: '{key}'")

    meta       = data["meta"]
    kpi        = data["kpi"]
    trend      = data.get("trend")
    breakdowns = data.get("breakdowns", [])
    point_a    = meta.get("point_a", "A")
    point_b    = meta.get("point_b", "B")

    meta.setdefault("generated", date.today().isoformat())

    if "findings" not in data or "recommendations" not in data:
        findings, recs = _build_findings(kpi, breakdowns, point_a, point_b)
        data.setdefault("findings", findings)
        data.setdefault("recommendations", recs)

    findings        = data["findings"]
    recommendations = data["recommendations"]

    # ── Open template ─────────────────────────────────────────────────────────
    prs    = Presentation(_TEMPLATE)
    slides = prs.slides

    # Template has exactly 6 slides: 0=title, 1=KPI, 2=trend, 3=bd0, 4=bd1, 5=findings

    _fill_slide0(slides[0], meta)
    _fill_slide1(slides[1], meta, kpi)
    _fill_slide2(slides[2], trend)

    # ── Fill breakdown slides (template has 2: slide 3 & 4) ──────────────────
    active_bds = breakdowns[:2]
    for i, bd in enumerate(active_bds):
        _fill_breakdown_slide(slides[3 + i], bd)

    # Blank out unused breakdown slides
    for i in range(len(active_bds), 2):
        blank_slide = slides[3 + i]
        for shape in blank_slide.shapes:
            if shape.has_text_frame:
                _set_text(shape, "")

    # ── Fill findings slide ───────────────────────────────────────────────────
    _fill_slide_findings(slides[5], meta, findings, recommendations)

    # ── Serialise to bytes ────────────────────────────────────────────────────
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
